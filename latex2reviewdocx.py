#!/usr/bin/env python3
"""Convert a LaTeX project to an editable review DOCX with pandoc."""

import argparse
import posixpath
import re
import shlex
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path


DEFAULT_PANDOC_MODULE = "pandoc/2.19.2"
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_CSL = SCRIPT_DIR / "simple_numeric.csl"
DEFAULT_FIGURE_FILTER = SCRIPT_DIR / "figure_caption_prefix.lua"
DEFAULT_EQUATION_FILTER = SCRIPT_DIR / "equation_numbering.lua"
DEFAULT_REFERENCE_DOC = SCRIPT_DIR / "reference_manuscript.docx"


class ValidationResult:
    def __init__(self):
        self.paragraphs = None
        self.tables = None
        self.inline_shapes = None
        self.track_revisions = False
        self.math_omml_count = 0
        self.media_count = 0


def run(cmd, dry_run=False, cwd=None):
    print("+ " + " ".join(shlex.quote(part) for part in cmd))
    if dry_run:
        return
    subprocess.run(cmd, check=True, cwd=cwd)


def remote_quote(path):
    return shlex.quote(path)


def validate_inputs(project_dir, main_tex, bib, csl, figure_filter=None, reference_doc=None, equation_filter=None):
    if not project_dir.is_dir():
        raise SystemExit(f"Project directory does not exist: {project_dir}")
    main_path = project_dir / main_tex
    if not main_path.is_file():
        raise SystemExit(f"Main TeX file does not exist: {main_path}")
    if bib and not (project_dir / bib).is_file():
        raise SystemExit(f"BibTeX file does not exist: {project_dir / bib}")
    if not csl.is_file():
        raise SystemExit(f"CSL file does not exist: {csl}")
    if figure_filter and not figure_filter.is_file():
        raise SystemExit(f"Figure caption filter does not exist: {figure_filter}")
    if reference_doc and not reference_doc.is_file():
        raise SystemExit(f"Reference DOCX does not exist: {reference_doc}")
    if equation_filter and not equation_filter.is_file():
        raise SystemExit(f"Equation numbering filter does not exist: {equation_filter}")
    missing_figures = find_missing_graphics(main_path, project_dir)
    if missing_figures:
        print("Warning: possible missing graphics:")
        for figure in missing_figures:
            print(f"  - {figure}")


def find_project_root(extract_dir):
    tex_files = [p for p in extract_dir.rglob("*.tex") if p.is_file()]
    if not tex_files:
        raise SystemExit("No .tex files found in the zip package.")
    roots = sorted({p.parent for p in tex_files}, key=lambda p: len(p.parts))
    if len(roots) == 1:
        return roots[0]
    top_dirs = [p for p in extract_dir.iterdir() if p.is_dir()]
    if len(top_dirs) == 1:
        return top_dirs[0]
    return extract_dir


def looks_like_main_tex(path):
    text = path.read_text(errors="replace")
    return "\\documentclass" in text and "\\begin{document}" in text


def detect_main_tex(project_dir, requested=None):
    if requested:
        main_path = project_dir / requested
        if not main_path.is_file():
            raise SystemExit(f"Main TeX file does not exist: {main_path}")
        return requested

    candidates = []
    for tex in project_dir.rglob("*.tex"):
        if tex.is_file() and looks_like_main_tex(tex):
            rel = tex.relative_to(project_dir).as_posix()
            score = 0
            name = tex.name.lower()
            if name in ("main.tex", "ms.tex", "manuscript.tex", "paper.tex", "article.tex"):
                score -= 10
            score += len(tex.relative_to(project_dir).parts)
            candidates.append((score, rel))

    if not candidates:
        raise SystemExit("Could not detect the main .tex file. Use --main-tex main.tex.")
    candidates.sort()
    return candidates[0][1]


def detect_bib(project_dir, requested=None):
    if requested:
        bib_path = project_dir / requested
        if not bib_path.is_file():
            raise SystemExit(f"BibTeX file does not exist: {bib_path}")
        return requested
    bibs = sorted([p.relative_to(project_dir).as_posix() for p in project_dir.rglob("*.bib") if p.is_file()])
    if len(bibs) == 1:
        return bibs[0]
    if len(bibs) > 1:
        preferred = [b for b in bibs if Path(b).name.lower() in ("references.bib", "refs.bib", "sample.bib")]
        if preferred:
            return preferred[0]
        print("Warning: multiple .bib files found; no bibliography was auto-selected.")
        for bib in bibs:
            print(f"  - {bib}")
    return None


def extract_zip(zip_path, extract_dir):
    if not zip_path.is_file():
        raise SystemExit(f"Zip package does not exist: {zip_path}")
    if zip_path.suffix.lower() != ".zip":
        raise SystemExit(f"Expected a .zip file: {zip_path}")
    with zipfile.ZipFile(zip_path) as z:
        z.extractall(extract_dir)


def strip_caption_short_options(text):
    out = []
    i = 0
    changed = 0
    token = r"\caption"
    while i < len(text):
        if not text.startswith(token, i):
            out.append(text[i])
            i += 1
            continue

        out.append(token)
        i += len(token)
        while i < len(text) and text[i].isspace():
            out.append(text[i])
            i += 1
        if i >= len(text) or text[i] != "[":
            continue

        depth = 1
        j = i + 1
        while j < len(text) and depth:
            if text[j] == "\\":
                j += 2
                continue
            if text[j] == "[":
                depth += 1
            elif text[j] == "]":
                depth -= 1
            j += 1

        if depth == 0:
            i = j
            changed += 1
        else:
            out.append("[")
            i += 1
    return "".join(out), changed


def normalize_includegraphics_options(text):
    pattern = re.compile(r"\\includegraphics\s*\[([^\]]*)\]\s*\{([^}]+)\}")
    changed = 0

    def replace(match):
        nonlocal changed
        options = match.group(1)
        graphic = match.group(2)
        option_names = {part.split("=", 1)[0].strip().lower() for part in options.split(",")}
        if "height" in option_names and "width" not in option_names:
            changed += 1
            return rf"\includegraphics[width=0.95\linewidth]{{{graphic}}}"
        return match.group(0)

    return pattern.sub(replace, text), changed


def read_latex_group(text, open_index):
    if open_index >= len(text) or text[open_index] != "{":
        return None
    depth = 1
    i = open_index + 1
    while i < len(text) and depth:
        if text[i] == "\\":
            i += 2
            continue
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
        i += 1
    if depth:
        return None
    return text[open_index + 1 : i - 1], i


def move_leading_caption_labels(text):
    out = []
    i = 0
    changed = 0
    token = r"\caption"
    label_pattern = re.compile(r"^\s*(\\label\{[^}]+\})(.*)$", re.DOTALL)
    while i < len(text):
        if not text.startswith(token, i):
            out.append(text[i])
            i += 1
            continue

        start = i
        i += len(token)
        j = i
        while j < len(text) and text[j].isspace():
            j += 1
        if j >= len(text) or text[j] != "{":
            out.append(text[start:i])
            continue

        group = read_latex_group(text, j)
        if not group:
            out.append(text[start:i])
            continue
        body, end = group
        match = label_pattern.match(body)
        if match:
            label = match.group(1)
            caption_body = match.group(2).strip()
            out.append(r"\caption{" + caption_body + "}\n" + label)
            changed += 1
        else:
            out.append(text[start:end])
        i = end
    return "".join(out), changed


def find_latex_command_group(text, command):
    idx = text.find(command)
    if idx < 0:
        return None
    i = idx + len(command)
    while i < len(text) and text[i].isspace():
        i += 1
    if i >= len(text) or text[i] != "{":
        return None
    group = read_latex_group(text, i)
    if not group:
        return None
    body, end = group
    return idx, body, end


def simplify_figure_environments(text):
    pattern = re.compile(r"\\begin\{figure\}(\[[^\]]*\])?(.*?)\\end\{figure\}", re.DOTALL)
    changed = 0

    def replace(match):
        nonlocal changed
        env = match.group(0)
        body = match.group(2)
        graphic_match = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{[^}]+\}", body)
        caption = find_latex_command_group(body, r"\caption")
        if not graphic_match or not caption:
            return env
        caption_body = caption[1].strip()
        label_match = re.search(r"\\label\{[^}]+\}", body)
        label = label_match.group(0) if label_match else ""
        pieces = [
            r"\begin{figure}",
            r"\centering",
            graphic_match.group(0),
            r"\caption{" + caption_body + "}",
        ]
        if label:
            pieces.append(label)
        pieces.append(r"\end{figure}")
        changed += 1
        return "\n".join(pieces)

    return pattern.sub(replace, text), changed


def preprocess_latex_for_pandoc(main_path):
    text = main_path.read_text(encoding="utf-8", errors="replace")
    text, captions_changed = strip_caption_short_options(text)
    text, caption_labels_changed = move_leading_caption_labels(text)
    text, figures_changed = simplify_figure_environments(text)
    text, graphics_changed = normalize_includegraphics_options(text)
    if captions_changed or caption_labels_changed or figures_changed or graphics_changed:
        main_path.write_text(text, encoding="utf-8")
        print(
            "Pandoc compatibility cleanup: "
            f"{captions_changed} caption short option(s), "
            f"{caption_labels_changed} caption label(s), "
            f"{figures_changed} figure environment(s), "
            f"{graphics_changed} height-only image option(s)."
        )


def find_missing_graphics(main_path, project_dir):
    text = main_path.read_text(errors="replace")
    matches = re.findall(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", text)
    missing = []
    extensions = ["", ".pdf", ".png", ".jpg", ".jpeg", ".eps"]
    for raw in matches:
        candidate = raw.strip()
        if "://" in candidate:
            continue
        if not any((project_dir / f"{candidate}{ext}").exists() for ext in extensions):
            missing.append(candidate)
    return missing


def enable_track_changes(src_docx, out_docx):
    with tempfile.TemporaryDirectory() as td_raw:
        td = Path(td_raw)
        with zipfile.ZipFile(src_docx) as z:
            z.extractall(td)

        settings = td / "word/settings.xml"
        if settings.exists():
            xml = settings.read_text(encoding="utf-8")
        else:
            settings.parent.mkdir(parents=True, exist_ok=True)
            xml = (
                '<w:settings '
                'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '</w:settings>'
            )

        if "<w:trackRevisions" not in xml:
            xml = xml.replace("</w:settings>", "<w:trackRevisions/></w:settings>")
        settings.write_text(xml, encoding="utf-8")

        with zipfile.ZipFile(out_docx, "w", zipfile.ZIP_DEFLATED) as z:
            for p in td.rglob("*"):
                if p.is_file():
                    z.write(p, p.relative_to(td).as_posix())


def ensure_child(parent, tag, first=False):
    child = parent.find(tag)
    if child is None:
        child = ET.Element(tag)
        if first:
            parent.insert(0, child)
        else:
            parent.append(child)
    return child


def polish_docx(path):
    table_borders = (
        '<w:tblBorders>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
        '</w:tblBorders>'
    )
    table_center = '<w:jc w:val="center"/>'
    caption_center = '<w:jc w:val="center"/>'

    def patch_table_caption(match):
        nonlocal table_count
        paragraph = match.group(0)
        text_match = re.search(r'<w:t([^>]*)>(.*?)</w:t>', paragraph, flags=re.S)
        if not text_match:
            return paragraph
        visible = re.sub(r'<[^>]+>', '', text_match.group(2)).strip()
        if not visible:
            return paragraph
        table_count += 1
        if not re.match(r'Table\s+\d+\.', visible):
            replacement = '<w:t%s>Table %d. %s</w:t>' % (
                text_match.group(1),
                table_count,
                text_match.group(2),
            )
            paragraph = paragraph[:text_match.start()] + replacement + paragraph[text_match.end():]
        paragraph = re.sub(r'<w:jc\s+[^>]*/>', '', paragraph)
        paragraph = paragraph.replace('</w:pPr>', caption_center + '</w:pPr>', 1)
        return paragraph

    def patch_table_properties(match):
        tblpr = match.group(0)
        tblpr = re.sub(r'<w:jc\s+[^>]*/>', '', tblpr)
        tblpr = re.sub(r'<w:tblBorders>.*?</w:tblBorders>', '', tblpr, flags=re.S)
        tblw = re.search(r'<w:tblW\b[^>]*/>', tblpr)
        if tblw:
            insert_at = tblw.end()
            return tblpr[:insert_at] + table_center + table_borders + tblpr[insert_at:]
        return tblpr.replace('</w:tblPr>', table_center + table_borders + '</w:tblPr>', 1)

    def patch_caption_style(styles_xml):
        style_match = re.search(
            r'(<w:style\b[^>]*w:styleId="Caption"[^>]*>)(.*?)(</w:style>)',
            styles_xml,
            flags=re.S,
        )
        if not style_match:
            return styles_xml
        body = style_match.group(2)
        rpr_match = re.search(r'<w:rPr>(.*?)</w:rPr>', body, flags=re.S)
        additions = '<w:i/><w:color w:val="000000"/><w:sz w:val="20"/><w:szCs w:val="20"/>'
        if rpr_match:
            rpr = rpr_match.group(0)
            rpr = re.sub(r'<w:color\b[^>]*/>', '', rpr)
            rpr = re.sub(r'<w:sz\b[^>]*/>', '', rpr)
            rpr = re.sub(r'<w:szCs\b[^>]*/>', '', rpr)
            if '<w:i' not in rpr:
                rpr = rpr.replace('</w:rPr>', '<w:i/></w:rPr>')
            rpr = rpr.replace('</w:rPr>', '<w:color w:val="000000"/><w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr>')
            body = body[:rpr_match.start()] + rpr + body[rpr_match.end():]
        else:
            body = '<w:rPr>' + additions + '</w:rPr>' + body
        return styles_xml[:style_match.start(2)] + body + styles_xml[style_match.end(2):]

    def pdf_box_ratio(pdf_path):
        data = pdf_path.read_bytes().decode("latin-1", errors="ignore")

        def read_box(name):
            match = re.search(
                rf"/{name}\s*\[\s*([-+]?\d*\.?\d+)\s+([-+]?\d*\.?\d+)\s+"
                rf"([-+]?\d*\.?\d+)\s+([-+]?\d*\.?\d+)\s*\]",
                data,
            )
            if not match:
                return None
            x1, y1, x2, y2 = (float(v) for v in match.groups())
            width = abs(x2 - x1)
            height = abs(y2 - y1)
            if width <= 0 or height <= 0:
                return None
            return width / height

        crop_ratio = read_box("CropBox")
        media_ratio = read_box("MediaBox")
        if not crop_ratio or not media_ratio:
            return None
        if abs(crop_ratio - media_ratio) / crop_ratio < 0.03:
            return None
        return crop_ratio

    def patch_pdf_image_extents(docxml, rels_xml, media_dir):
        rid_to_target = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels_xml))
        fixed = 0

        def patch_drawing(match):
            nonlocal fixed
            drawing = match.group(0)
            rid_match = re.search(r'<a:blip\s+r:embed="([^"]+)"', drawing)
            extent_match = re.search(r'<wp:extent\s+cx="(\d+)"\s+cy="(\d+)"\s*/>', drawing)
            if not rid_match or not extent_match:
                return drawing
            target = rid_to_target.get(rid_match.group(1), "")
            if not target.lower().endswith(".pdf"):
                return drawing
            pdf_path = media_dir / Path(target).name
            if not pdf_path.is_file():
                return drawing
            ratio = pdf_box_ratio(pdf_path)
            if not ratio:
                return drawing
            cx = int(extent_match.group(1))
            cy = int(extent_match.group(2))
            new_cy = max(1, round(cx / ratio))
            if abs(new_cy - cy) / max(cy, 1) < 0.03:
                return drawing
            drawing = re.sub(
                r'<wp:extent\s+cx="\d+"\s+cy="\d+"\s*/>',
                f'<wp:extent cx="{cx}" cy="{new_cy}"/>',
                drawing,
                count=1,
            )
            drawing = re.sub(
                r'<a:ext\s+cx="\d+"\s+cy="\d+"\s*/>',
                f'<a:ext cx="{cx}" cy="{new_cy}"/>',
                drawing,
                count=1,
            )
            fixed += 1
            return drawing

        docxml = re.sub(r'<w:drawing>.*?</w:drawing>', patch_drawing, docxml, flags=re.S)
        if fixed:
            print(f"DOCX PDF image aspect cleanup: {fixed} image(s).")
        return docxml

    with tempfile.TemporaryDirectory() as td_raw:
        td = Path(td_raw)
        with zipfile.ZipFile(path) as z:
            z.extractall(td)

        docxml_path = td / "word/document.xml"
        docxml = docxml_path.read_text(encoding="utf-8")
        rels_path = td / "word/_rels/document.xml.rels"
        if rels_path.exists():
            docxml = patch_pdf_image_extents(
                docxml,
                rels_path.read_text(encoding="utf-8"),
                td / "word/media",
            )
        table_count = 0
        docxml = re.sub(
            r'<w:p><w:pPr><w:pStyle\s+w:val="TableCaption"\s*/>.*?</w:p>',
            patch_table_caption,
            docxml,
            flags=re.S,
        )
        docxml = re.sub(r'<w:tblPr>.*?</w:tblPr>', patch_table_properties, docxml, flags=re.S)
        docxml = re.sub(
            r'(<w:pgMar\b[^>]*w:top=")[^"]*("[^>]*w:bottom=")[^"]*("[^>]*w:left=")[^"]*("[^>]*w:right=")[^"]*(")',
            r'\g<1>1440\g<2>1440\g<3>1440\g<4>1440\g<5>',
            docxml,
        )
        docxml_path.write_text(docxml, encoding="utf-8")

        styles_path = td / "word/styles.xml"
        if styles_path.exists():
            styles_xml = styles_path.read_text(encoding="utf-8")
            styles_path.write_text(patch_caption_style(styles_xml), encoding="utf-8")

        tmp_out = path.with_name(path.stem + "_polished_tmp.docx")
        with zipfile.ZipFile(tmp_out, "w", zipfile.ZIP_DEFLATED) as z:
            for p in td.rglob("*"):
                if p.is_file():
                    z.write(p, p.relative_to(td).as_posix())
        tmp_out.replace(path)


def validate_docx(path):
    result = ValidationResult()
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            raise SystemExit(f"DOCX zip integrity check failed at: {bad}")
        docxml = z.read("word/document.xml").decode("utf-8", errors="replace")
        settings = z.read("word/settings.xml").decode("utf-8", errors="replace")
        result.track_revisions = "<w:trackRevisions" in settings
        result.math_omml_count = docxml.count("<m:oMath")
        result.media_count = len([n for n in z.namelist() if n.startswith("word/media/")])

    try:
        from docx import Document  # type: ignore

        doc = Document(path)
        result.paragraphs = len(doc.paragraphs)
        result.tables = len(doc.tables)
        result.inline_shapes = len(doc.inline_shapes)
    except Exception as exc:
        print(f"Warning: python-docx structural counts skipped: {exc}")
    return result


def print_validation(result):
    if result.paragraphs is not None:
        print(f"paragraphs: {result.paragraphs}")
        print(f"tables: {result.tables}")
        print(f"inline_shapes: {result.inline_shapes}")
    print(f"trackRevisions: {result.track_revisions}")
    print(f"math OMML count: {result.math_omml_count}")
    print(f"media count: {result.media_count}")
    print("zip integrity: OK")


def build_remote_command(args, remote_docx, remote_csl_name):
    parts = [
        f"cd {remote_quote(args.remote_project_dir)}",
        f"module load {shlex.quote(args.pandoc_module)}",
        "pandoc",
        shlex.quote(args.main_tex),
        "--citeproc",
        f"--csl={shlex.quote(remote_csl_name)}",
        "--resource-path=.",
        "-s",
        f"-o {remote_quote(remote_docx)}",
    ]
    if args.bib:
        parts.insert(-3, f"--bibliography={shlex.quote(args.bib)}")
    if getattr(args, "figure_prefix", False):
        parts.insert(-3, f"--lua-filter={shlex.quote(Path(args.figure_filter).name)}")
    if getattr(args, "equation_numbers", False):
        parts.insert(-3, f"--lua-filter={shlex.quote(Path(args.equation_filter).name)}")
    if getattr(args, "use_reference_doc", False):
        parts.insert(-3, f"--reference-doc={shlex.quote(Path(args.reference_doc).name)}")
    if args.extra_pandoc_arg:
        parts[-3:-3] = [shlex.quote(arg) for arg in args.extra_pandoc_arg]
    return " ".join(parts[:2]) + " && " + " ".join(parts[2:])


def build_pandoc_cmd(args, csl, output):
    cmd = [
        "pandoc",
        args.main_tex,
        "--citeproc",
        f"--csl={csl}",
        "--resource-path=.",
        "-s",
        "-o",
        str(output),
    ]
    if args.bib:
        cmd.insert(-4, f"--bibliography={args.bib}")
    if getattr(args, "figure_prefix", False):
        cmd.insert(-4, f"--lua-filter={args.figure_filter}")
    if getattr(args, "equation_numbers", False):
        cmd.insert(-4, f"--lua-filter={args.equation_filter}")
    if getattr(args, "use_reference_doc", False):
        cmd.insert(-4, f"--reference-doc={args.reference_doc}")
    if args.extra_pandoc_arg:
        cmd[-4:-4] = args.extra_pandoc_arg
    return cmd


def convert_local(args):
    project_dir = Path(args.project_dir).expanduser().resolve()
    output = Path(args.output).expanduser().resolve()
    csl = Path(args.numeric_csl).expanduser().resolve()
    figure_filter = Path(args.figure_filter).expanduser().resolve() if args.figure_prefix else None
    reference_doc = Path(args.reference_doc).expanduser().resolve() if args.use_reference_doc else None
    equation_filter = Path(args.equation_filter).expanduser().resolve() if args.equation_numbers else None
    if figure_filter:
        args.figure_filter = str(figure_filter)
    if reference_doc:
        args.reference_doc = str(reference_doc)
    if equation_filter:
        args.equation_filter = str(equation_filter)
    validate_inputs(project_dir, args.main_tex, args.bib, csl, figure_filter, reference_doc, equation_filter)
    output.parent.mkdir(parents=True, exist_ok=True)

    raw_output = output
    if args.track_changes:
        raw_output = output.with_name(output.stem + "_raw_pandoc.docx")

    cmd = build_pandoc_cmd(args, csl, raw_output)
    if args.load_module:
        cmd = ["bash", "-lc", f"module load {shlex.quote(args.pandoc_module)} && " + " ".join(shlex.quote(x) for x in cmd)]
    print(f"Working directory: {project_dir}")
    run(cmd, dry_run=args.dry_run, cwd=project_dir)

    if args.dry_run:
        print(f"Dry run complete. Planned output: {output}")
        return

    if args.track_changes:
        enable_track_changes(raw_output, output)
        if not args.keep_raw:
            if raw_output.exists():
                raw_output.unlink()
    if args.layout_polish:
        polish_docx(output)

    result = validate_docx(output)
    print_validation(result)
    print(f"Final DOCX: {output}")


def convert_zip(args):
    zip_path = Path(args.zip_path).expanduser().resolve()
    output = Path(args.output).expanduser().resolve()
    csl = Path(args.numeric_csl).expanduser().resolve()
    if not csl.is_file():
        raise SystemExit(f"CSL file does not exist: {csl}")
    figure_filter = Path(args.figure_filter).expanduser().resolve() if args.figure_prefix else None
    reference_doc = Path(args.reference_doc).expanduser().resolve() if args.use_reference_doc else None
    equation_filter = Path(args.equation_filter).expanduser().resolve() if args.equation_numbers else None
    if figure_filter:
        args.figure_filter = str(figure_filter)
        if not figure_filter.is_file():
            raise SystemExit(f"Figure caption filter does not exist: {figure_filter}")
    if reference_doc:
        args.reference_doc = str(reference_doc)
        if not reference_doc.is_file():
            raise SystemExit(f"Reference DOCX does not exist: {reference_doc}")
    if equation_filter:
        args.equation_filter = str(equation_filter)
        if not equation_filter.is_file():
            raise SystemExit(f"Equation numbering filter does not exist: {equation_filter}")
    output.parent.mkdir(parents=True, exist_ok=True)

    work_root = None
    if args.work_dir:
        work_root = Path(args.work_dir).expanduser().resolve()
        if work_root.exists():
            if not args.reuse_work_dir:
                raise SystemExit(f"Work directory already exists: {work_root}")
        else:
            work_root.mkdir(parents=True)
    else:
        work_root = Path(tempfile.mkdtemp(prefix="latex2reviewdocx_"))

    try:
        extract_dir = work_root / "source"
        extract_dir.mkdir(parents=True, exist_ok=True)
        extract_zip(zip_path, extract_dir)
        project_dir = find_project_root(extract_dir)
        main_tex = detect_main_tex(project_dir, args.main_tex)
        bib = detect_bib(project_dir, args.bib)

        print(f"Zip package: {zip_path}")
        print(f"Project directory: {project_dir}")
        print(f"Main TeX: {main_tex}")
        if bib:
            print(f"Bibliography: {bib}")
        else:
            print("Bibliography: none")

        preprocess_latex_for_pandoc(project_dir / main_tex)

        local_args = argparse.Namespace(
            project_dir=str(project_dir),
            main_tex=main_tex,
            bib=bib,
            output=str(output),
            numeric_csl=str(csl),
            pandoc_module=args.pandoc_module,
            track_changes=args.track_changes,
            keep_raw=args.keep_raw,
            dry_run=args.dry_run,
            extra_pandoc_arg=args.extra_pandoc_arg,
            load_module=args.load_module,
            figure_prefix=args.figure_prefix,
            figure_filter=args.figure_filter,
            equation_numbers=args.equation_numbers,
            equation_filter=args.equation_filter,
            use_reference_doc=args.use_reference_doc,
            reference_doc=args.reference_doc,
            layout_polish=args.layout_polish,
        )
        convert_local(local_args)
    finally:
        if args.work_dir:
            print(f"Work directory kept: {work_root}")
        elif not args.dry_run:
            shutil.rmtree(str(work_root), ignore_errors=True)


def convert_remote(args):
    project_dir = Path(args.project_dir).expanduser().resolve()
    output = Path(args.output).expanduser().resolve()
    csl = Path(args.numeric_csl).expanduser().resolve()
    figure_filter = Path(args.figure_filter).expanduser().resolve() if args.figure_prefix else None
    reference_doc = Path(args.reference_doc).expanduser().resolve() if args.use_reference_doc else None
    equation_filter = Path(args.equation_filter).expanduser().resolve() if args.equation_numbers else None
    if figure_filter:
        args.figure_filter = str(figure_filter)
    if reference_doc:
        args.reference_doc = str(reference_doc)
    if equation_filter:
        args.equation_filter = str(equation_filter)
    validate_inputs(project_dir, args.main_tex, args.bib, csl, figure_filter, reference_doc, equation_filter)
    output.parent.mkdir(parents=True, exist_ok=True)

    remote_project = f"{args.remote}:{args.remote_project_dir.rstrip('/')}/"
    run(
        [
            "rsync",
            "-av",
            "--exclude",
            ".git",
            "--exclude",
            ".DS_Store",
            f"{project_dir}/",
            remote_project,
        ],
        dry_run=args.dry_run,
    )

    remote_csl_name = Path(csl).name
    run(
        [
            "rsync",
            "-av",
            str(csl),
            f"{args.remote}:{posixpath.join(args.remote_project_dir, remote_csl_name)}",
        ],
        dry_run=args.dry_run,
    )
    if args.figure_prefix:
        run(
            [
                "rsync",
                "-av",
                str(figure_filter),
                f"{args.remote}:{posixpath.join(args.remote_project_dir, Path(args.figure_filter).name)}",
            ],
            dry_run=args.dry_run,
        )
    if args.equation_numbers:
        run(
            [
                "rsync",
                "-av",
                str(equation_filter),
                f"{args.remote}:{posixpath.join(args.remote_project_dir, Path(args.equation_filter).name)}",
            ],
            dry_run=args.dry_run,
        )
    if args.use_reference_doc:
        run(
            [
                "rsync",
                "-av",
                str(reference_doc),
                f"{args.remote}:{posixpath.join(args.remote_project_dir, Path(args.reference_doc).name)}",
            ],
            dry_run=args.dry_run,
        )

    remote_base = output.with_suffix("").name + "_pandoc.docx"
    remote_docx_path = posixpath.join(args.remote_project_dir, remote_base)
    remote_command = build_remote_command(args, remote_docx_path, remote_csl_name)
    run(["ssh", args.remote, remote_command], dry_run=args.dry_run)

    pulled = output
    if args.track_changes:
        pulled = output.with_name(output.stem + "_raw_pandoc.docx")
    run(["rsync", "-av", f"{args.remote}:{remote_docx_path}", str(pulled)], dry_run=args.dry_run)

    if args.dry_run:
        print(f"Dry run complete. Planned output: {output}")
        return

    if args.track_changes:
        enable_track_changes(pulled, output)
        if not args.keep_raw:
            if pulled.exists():
                pulled.unlink()
    if args.layout_polish:
        polish_docx(output)

    result = validate_docx(output)
    print_validation(result)
    print(f"Final DOCX: {output}")


def parse_args(argv):
    parser = argparse.ArgumentParser(
        description="Convert a LaTeX manuscript to editable DOCX with pandoc.",
    )
    subparsers = parser.add_subparsers(dest="mode")

    def add_shared_options(subparser):
        subparser.add_argument("--numeric-csl", default=str(DEFAULT_CSL), help="Numeric CSL file.")
        subparser.add_argument(
            "--figure-filter",
            default=str(DEFAULT_FIGURE_FILTER),
            help="Lua filter used to add Fig. N. prefixes to figure captions.",
        )
        subparser.add_argument(
            "--no-figure-prefix",
            dest="figure_prefix",
            action="store_false",
            help="Do not add Fig. N. prefixes to figure captions.",
        )
        subparser.set_defaults(figure_prefix=True)
        subparser.add_argument(
            "--equation-filter",
            default=str(DEFAULT_EQUATION_FILTER),
            help="Lua filter used to add (N) labels to display equations.",
        )
        subparser.add_argument(
            "--no-equation-numbers",
            dest="equation_numbers",
            action="store_false",
            help="Do not add (N) labels to display equations.",
        )
        subparser.set_defaults(equation_numbers=True)
        subparser.add_argument(
            "--reference-doc",
            default=str(DEFAULT_REFERENCE_DOC),
            help="Word reference template used to improve manuscript-like formatting.",
        )
        subparser.add_argument(
            "--no-reference-doc",
            dest="use_reference_doc",
            action="store_false",
            help="Do not use the bundled Word reference template.",
        )
        subparser.set_defaults(use_reference_doc=True)
        subparser.add_argument(
            "--no-layout-polish",
            dest="layout_polish",
            action="store_false",
            help="Skip DOCX post-processing for captions and page margins.",
        )
        subparser.set_defaults(layout_polish=True)
        subparser.add_argument("--pandoc-module", default=DEFAULT_PANDOC_MODULE, help="Pandoc environment module.")
        subparser.add_argument("--track-changes", action="store_true", help="Enable Word Track Changes mode.")
        subparser.add_argument("--keep-raw", action="store_true", help="Keep the unpatched pandoc DOCX.")
        subparser.add_argument("--dry-run", action="store_true", help="Print commands without executing them.")
        subparser.add_argument(
            "--extra-pandoc-arg",
            action="append",
            help="Additional pandoc argument. Repeat for multiple arguments.",
        )

    def add_common_options(subparser):
        subparser.add_argument("--project-dir", required=True, help="LaTeX project directory.")
        subparser.add_argument("--main-tex", required=True, help="Main .tex file relative to project dir.")
        subparser.add_argument("--bib", help="BibTeX file relative to project dir.")
        subparser.add_argument("--output", required=True, help="Output .docx path.")
        add_shared_options(subparser)

    zip_parser = subparsers.add_parser("zip", help="Convert an Overleaf zip package.")
    zip_parser.add_argument("zip_path", help="Path to the Overleaf .zip package.")
    zip_parser.add_argument("output", help="Output .docx path.")
    zip_parser.add_argument("--main-tex", help="Main .tex file relative to the detected project root.")
    zip_parser.add_argument("--bib", help="BibTeX file relative to the detected project root.")
    zip_parser.add_argument("--work-dir", help="Optional extraction/work directory to keep.")
    zip_parser.add_argument("--reuse-work-dir", action="store_true", help="Allow using an existing work directory.")
    zip_parser.add_argument(
        "--load-module",
        action="store_true",
        help="Load the pandoc environment module before running pandoc.",
    )
    add_shared_options(zip_parser)
    zip_parser.set_defaults(func=convert_zip)

    local_parser = subparsers.add_parser("local", help="Run pandoc directly on this machine.")
    add_common_options(local_parser)
    local_parser.add_argument(
        "--load-module",
        action="store_true",
        help="Load the pandoc environment module before running pandoc.",
    )
    local_parser.set_defaults(func=convert_local)

    remote_parser = subparsers.add_parser("remote", help="Sync to a remote host, convert there, and pull DOCX back.")
    add_common_options(remote_parser)
    remote_parser.add_argument("--remote", required=True, help="SSH remote, for example user@cluster.example.edu.")
    remote_parser.add_argument(
        "--remote-project-dir",
        required=True,
        help="Remote working directory where the project will be synced.",
    )
    remote_parser.set_defaults(func=convert_remote)
    args = parser.parse_args(argv)
    if not hasattr(args, "func"):
        parser.print_help()
        raise SystemExit(2)
    return args


def main(argv=None):
    args = parse_args(argv or sys.argv[1:])
    try:
        args.func(args)
    except subprocess.CalledProcessError as exc:
        print(f"Command failed with exit code {exc.returncode}", file=sys.stderr)
        return exc.returncode
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
